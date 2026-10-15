"""
Ingest-stap: documenten inlezen -> chunks -> embeddings -> index.json.

Ondersteunt .md, .txt EN .docx (Word), zodat je je eigen kennisbank-documenten
rechtstreeks kunt gebruiken. Draai opnieuw na het wijzigen/toevoegen van documenten:

    python src/ingest.py
"""

from __future__ import annotations
import json
import sys
from pathlib import Path

from config import DOCS_DIR, INDEX_PATH, CHUNK_SIZE, CHUNK_OVERLAP
from lmclient import embed


def _read_docx(path: Path) -> str:
    """Leest tekst (alinea's + tabellen) uit een Word-bestand via python-docx."""
    try:
        import docx
    except ImportError:
        print(f"  LET OP: python-docx ontbreekt, {path.name} overgeslagen. "
              f"Installeer met: pip install python-docx")
        return ""
    doc = docx.Document(str(path))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return "\n\n".join(parts)


def read_documents(docs_dir: Path) -> list[tuple[str, str]]:
    """Leest alle .md/.txt/.docx-bestanden. Geeft (bestandsnaam, inhoud) terug."""
    docs = []
    for path in sorted(docs_dir.glob("*")):
        suffix = path.suffix.lower()
        if suffix in (".md", ".txt"):
            docs.append((path.name, path.read_text(encoding="utf-8")))
        elif suffix == ".docx" and not path.name.startswith("~$"):
            text = _read_docx(path)
            if text:
                docs.append((path.name, text))
    return docs


def chunk_text(text: str, size: int, overlap: int) -> list[str]:
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            break_point = text.rfind("\n\n", start, end)
            if break_point > start + size // 2:
                end = break_point
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def build_index() -> None:
    docs = read_documents(DOCS_DIR)
    if not docs:
        print(f"Geen documenten gevonden in {DOCS_DIR}. Zet er .md/.txt/.docx in.")
        sys.exit(1)

    records = []
    for filename, content in docs:
        chunks = chunk_text(content, CHUNK_SIZE, CHUNK_OVERLAP)
        print(f"  {filename}: {len(chunks)} chunk(s)")
        for i, chunk in enumerate(chunks):
            records.append({
                "source": filename,
                "chunk_id": i,
                "text": chunk,
                "embedding": embed(chunk),
            })

    INDEX_PATH.write_text(json.dumps(records, ensure_ascii=False), encoding="utf-8")
    print(f"\nKlaar. {len(records)} chunks geïndexeerd -> {INDEX_PATH.name}")


if __name__ == "__main__":
    print(f"Documenten inlezen uit: {DOCS_DIR}")
    build_index()
